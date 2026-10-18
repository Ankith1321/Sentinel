import hashlib
import io
import logging
from pathlib import PurePosixPath

from docx import Document
from pypdf import PdfReader

from backend.app import store
from ingestion.chunker import ChunkRecord, chunk_record
from ingestion.config import UPLOAD_PATH_PREFIX
from ingestion.loader import stable_source_id
from ingestion.records import SourceRecord
from ingestion.retriever import get_collection, get_keyword_index, get_model

logger = logging.getLogger(__name__)

EMBED_BATCH_SIZE = 64
LANGUAGE_BY_SUFFIX = {
    ".py": "python",
    ".md": "markdown",
    ".mdx": "markdown",
    ".js": "javascript",
    ".jsx": "javascript",
    ".ts": "typescript",
    ".tsx": "typescript",
    ".txt": "text",
    ".pdf": "text",
    ".docx": "text",
}
SUPPORTED_UPLOAD_SUFFIXES = frozenset(LANGUAGE_BY_SUFFIX)


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = PurePosixPath(safe_filename(filename)).suffix.lower()

    if suffix not in SUPPORTED_UPLOAD_SUFFIXES:
        raise UnsupportedDocument(f"Unsupported file type: {suffix or filename}")

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in document.paragraphs if p.text.strip())

    return data.decode("utf-8", errors="replace")


def safe_filename(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).name


def upload_path(filename: str) -> str:
    return f"{UPLOAD_PATH_PREFIX}{safe_filename(filename)}"


def embed_and_upsert(chunks: list[ChunkRecord]) -> None:
    collection = get_collection()
    model = get_model()

    for start in range(0, len(chunks), EMBED_BATCH_SIZE):
        batch = chunks[start:start + EMBED_BATCH_SIZE]
        texts = [c.content for c in batch]
        embeddings = model.encode(texts, batch_size=EMBED_BATCH_SIZE, normalize_embeddings=True).tolist()

        collection.upsert(
            ids=[c.chunk_id for c in batch],
            documents=texts,
            embeddings=embeddings,
            metadatas=[
                {
                    "source_id": c.source_id,
                    "source_type": c.source_type,
                    "path": c.path,
                    "language": c.language,
                    "title": c.title or "",
                    "chunk_index": c.chunk_index,
                }
                for c in batch
            ],
        )


def ingest_upload(job_id: str, filename: str, data: bytes) -> None:
    store.update_job(job_id, "processing")

    try:
        path = upload_path(filename)
        text = extract_text(filename, data).strip()
        if not text:
            raise UnsupportedDocument("No extractable text in file")

        content_hash = hashlib.sha256(data).hexdigest()
        existing = store.get_document_by_path(path)

        if existing and existing["content_hash"] == content_hash:
            store.update_job(job_id, "unchanged", document_id=existing["id"])
            return

        record = SourceRecord(
            source_id=stable_source_id(path),
            source_type="upload",
            path=path,
            language=LANGUAGE_BY_SUFFIX[PurePosixPath(safe_filename(filename)).suffix.lower()],
            title=safe_filename(filename),
            content=text,
            content_hash=content_hash,
        )
        chunks = chunk_record(record)

        get_collection().delete(where={"source_id": record.source_id})
        if chunks:
            embed_and_upsert(chunks)

        store.upsert_document(record.source_id, path, record.source_type, record.language, content_hash, len(chunks))
        get_keyword_index.cache_clear()
        get_keyword_index()

        store.update_job(job_id, "done", document_id=record.source_id)
        logger.info("ingested %s as %d chunks", path, len(chunks))
    except Exception as error:
        logger.exception("ingestion failed for %s", filename)
        store.update_job(job_id, "failed", error=str(error))


def delete_indexed_document(document_id: str) -> bool:
    if store.get_document(document_id) is None:
        return False

    get_collection().delete(where={"source_id": document_id})
    store.delete_document(document_id)
    get_keyword_index.cache_clear()
    return True
